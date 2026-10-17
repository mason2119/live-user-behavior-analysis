# -*- coding: utf-8 -*-
"""
问题2：高价值用户画像
方法：K-Means聚类 + Spearman相关性分析 + LinearSVC + 特征重要性 + 混淆矩阵
目标：识别高价值用户特征，构建付费预测模型
"""

import os
import sys
import warnings
import datetime
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from sklearn.model_selection import train_test_split, RandomizedSearchCV
from sklearn.preprocessing import StandardScaler, OneHotEncoder, LabelEncoder
from sklearn.compose import ColumnTransformer
from sklearn.pipeline import Pipeline
from sklearn.svm import LinearSVC
from sklearn.neural_network import MLPClassifier
from sklearn.cluster import KMeans
from sklearn.metrics import (confusion_matrix, classification_report,
                             roc_auc_score, accuracy_score, precision_score,
                             recall_score, f1_score, roc_curve,
                             silhouette_score)
from sklearn.impute import SimpleImputer
from scipy.stats import spearmanr
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

warnings.filterwarnings('ignore')

# ============================================================
# 一、配置
# ============================================================
_CODE_DIR = os.path.dirname(os.path.abspath(__file__))
_ROOT_DIR = os.path.dirname(_CODE_DIR)
DATA_DIR = os.path.join(_ROOT_DIR, 'data', '中文版')
OUTPUT_DIR = os.path.join(_CODE_DIR, 'output')
DOC_DIR = os.path.join(_ROOT_DIR, 'reports')
LOG_FILE = os.path.join(OUTPUT_DIR, '09_问题2_分析日志.txt')
DOC_FILE = os.path.join(DOC_DIR, '09_问题2_高价值用户画像分析报告.docx')

os.makedirs(OUTPUT_DIR, exist_ok=True)
os.environ['JOBLIB_TEMP_FOLDER'] = os.path.abspath('./joblib_temp')
os.makedirs(os.environ['JOBLIB_TEMP_FOLDER'], exist_ok=True)

# 避免 joblib/loky 在 Windows 下自动检测CPU核心时因编码错误崩溃
os.environ['LOKY_MAX_CPU_COUNT'] = '4'
os.environ['LOKY_PICKLER'] = 'pickle'

plt.rcParams['font.sans-serif'] = ['SimHei', 'Microsoft YaHei']
plt.rcParams['axes.unicode_minus'] = False


# ============================================================
# 二、日志工具
# ============================================================
class Tee:
    def __init__(self, filename):
        self.terminal = sys.stdout
        self.log = open(filename, 'w', encoding='utf-8')

    def write(self, message):
        self.terminal.write(message)
        self.log.write(message)
        self.flush()

    def flush(self):
        self.terminal.flush()
        self.log.flush()


# ============================================================
# 三、数据加载与特征工程
# ============================================================
def load_and_engineer_features():
    """加载数据并构建用户级特征"""
    print("=" * 80)
    print("正在加载数据并构建用户级特征...")
    print("=" * 80)

    # 1. 用户基础信息
    user_path = os.path.join(DATA_DIR, '用户基础信息.csv')
    df_user = pd.read_csv(user_path)
    print(f"用户基础信息表: {df_user.shape[0]:,} 行 × {df_user.shape[1]} 列")

    # 处理年龄异常值
    df_user.loc[(df_user['年龄'] < 10) | (df_user['年龄'] > 80), '年龄'] = np.nan

    # 2. 用户行为日表聚合特征
    behavior_path = os.path.join(DATA_DIR, '用户行为日表.csv')
    df_behavior = pd.read_csv(behavior_path)
    print(f"用户行为日表: {df_behavior.shape[0]:,} 行 × {df_behavior.shape[1]} 列")

    # 过滤观看时长负值
    df_behavior = df_behavior[df_behavior['总观看时长(秒)'] >= 0].copy()

    behavior_agg = df_behavior.groupby('用户ID').agg({
        '总观看时长(秒)': ['mean', 'sum', 'count'],
        '点赞次数': ['mean', 'sum'],
        '评论次数': ['mean', 'sum'],
        '关注次数': ['mean', 'sum'],
        '参与PK次数': ['mean', 'sum'],
        '送礼次数': ['mean', 'sum'],
        '当日送礼总价值': ['mean', 'sum'],
        '连续登录天数': 'max',
        '是否跨时区观看': 'mean',
    }).reset_index()

    # 扁平化多级列名
    behavior_agg.columns = ['用户ID'] + [
        f"{col[0]}_{col[1]}" if col[1] != '' else col[0]
        for col in behavior_agg.columns[1:]
    ]

    # 3. 送礼交易明细聚合特征（RFM）
    gift_path = os.path.join(DATA_DIR, '送礼交易明细.csv')
    df_gift = pd.read_csv(gift_path)
    print(f"送礼交易明细表: {df_gift.shape[0]:,} 行 × {df_gift.shape[1]} 列")

    # Recency: 最后一次送礼距今天数（假设最大日期为基准）
    df_gift['交易日期'] = pd.to_datetime(df_gift['交易日期'], format='mixed', dayfirst=False)
    max_date = df_gift['交易日期'].max()

    gift_agg = df_gift.groupby('送礼用户ID').agg({
        '交易日期': lambda x: (max_date - x.max()).days,  # Recency
        '交易ID': 'count',  # Frequency
        '礼物价值(美元)': ['sum', 'mean', 'max'],  # Monetary
        '是否PK场景送礼': 'mean',
        '是否首次送礼': 'max',
    }).reset_index()

    gift_agg.columns = ['用户ID', '最近送礼天数', '送礼频次',
                        '累计礼物价值', '平均礼物价值', '最大礼物价值',
                        'PK场景送礼比例', '是否首次送礼']

    # 4. 合并所有特征
    df = df_user.merge(behavior_agg, on='用户ID', how='left')
    df = df.merge(gift_agg, on='用户ID', how='left')

    # 5. 定义目标变量
    df['是否付费'] = (df['累计消费(美元)'] > 0).astype(int)
    # 高价值用户：在付费用户中取累计消费前20%
    paid_users = df[df['是否付费'] == 1].copy()
    spending_threshold = paid_users['累计消费(美元)'].quantile(0.8)
    df['是否高价值'] = (df['累计消费(美元)'] > spending_threshold).astype(int)

    print(f"\n合并后用户级数据: {df.shape[0]:,} 行 × {df.shape[1]} 列")
    print(f"付费用户比例: {df['是否付费'].mean():.2%}")
    print(f"高价值用户比例: {df['是否高价值'].mean():.2%}")

    return df


# ============================================================
# 四、K-Means聚类分析
# ============================================================
def kmeans_clustering(df):
    """对用户进行K-Means聚类"""
    print("\n" + "=" * 80)
    print("正在进行K-Means聚类分析...")
    print("=" * 80)

    # 聚类特征（数值型，需填充缺失）
    cluster_features = [
        '累计消费(美元)', '关注主播数', '社交强度得分',
        '总观看时长(秒)_mean', '点赞次数_mean', '评论次数_mean',
        '送礼频次', '累计礼物价值', 'VIP等级'
    ]
    cluster_features = [f for f in cluster_features if f in df.columns]

    df_cluster = df[cluster_features].copy()
    df_cluster = df_cluster.fillna(df_cluster.median())

    # 为加速聚类，采样最多5万条
    if len(df_cluster) > 50000:
        df_cluster_sample = df_cluster.sample(n=50000, random_state=42)
    else:
        df_cluster_sample = df_cluster.copy()
    print(f"用于聚类的样本数: {len(df_cluster_sample):,}")

    # 标准化
    scaler = StandardScaler()
    X_scaled = scaler.fit_transform(df_cluster_sample)

    # 选择最佳K值（肘部法则+轮廓系数）
    inertia = []
    silhouette_scores = []
    K_range = range(2, 6)

    for k in K_range:
        kmeans = KMeans(n_clusters=k, random_state=42, n_init=5)
        labels = kmeans.fit_predict(X_scaled)
        inertia.append(kmeans.inertia_)
        silhouette_scores.append(silhouette_score(X_scaled, labels))
        print(f"  K={k}: 轮廓系数={silhouette_scores[-1]:.4f}")

    # 选择轮廓系数最大的K
    best_k = list(K_range)[np.argmax(silhouette_scores)]
    print(f"\n最佳聚类数 K={best_k}")

    # 最终聚类（基于样本训练，再预测全部数据）
    final_kmeans = KMeans(n_clusters=best_k, random_state=42, n_init=5)
    final_kmeans.fit(X_scaled)
    X_scaled_all = scaler.transform(df_cluster)
    df['聚类标签'] = final_kmeans.predict(X_scaled_all)

    # 聚类画像
    cluster_profile = df.groupby('聚类标签').agg({
        '累计消费(美元)': 'mean',
        '是否付费': 'mean',
        '是否高价值': 'mean',
        '总观看时长(秒)_mean': 'mean',
        '点赞次数_mean': 'mean',
        '送礼频次': 'mean',
        'VIP等级': 'mean',
        '用户ID': 'count'
    }).round(4)
    cluster_profile.columns = ['平均累计消费', '付费率', '高价值率',
                               '平均观看时长', '平均点赞次数', '平均送礼频次',
                               '平均VIP等级', '用户数']

    print("\n聚类画像:")
    print(cluster_profile.to_string())

    # 保存聚类结果
    cluster_path = os.path.join(OUTPUT_DIR, '09_聚类画像结果.csv')
    cluster_profile.to_csv(cluster_path, encoding='utf-8-sig')
    print(f"\n聚类画像已保存: {cluster_path}")

    # 绘制肘部法则图
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(12, 5))
    ax1.plot(K_range, inertia, 'bo-')
    ax1.set_xlabel('聚类数 K')
    ax1.set_ylabel('SSE (Inertia)')
    ax1.set_title('K-Means 肘部法则')
    ax1.grid(True)

    ax2.plot(K_range, silhouette_scores, 'ro-')
    ax2.set_xlabel('聚类数 K')
    ax2.set_ylabel('轮廓系数')
    ax2.set_title('K-Means 轮廓系数')
    ax2.grid(True)

    plt.tight_layout()
    elbow_path = os.path.join(OUTPUT_DIR, '09_KMeans_聚类评估.png')
    plt.savefig(elbow_path, dpi=300, bbox_inches='tight')
    plt.close()
    print(f"聚类评估图已保存: {elbow_path}")

    return df, cluster_profile


# ============================================================
# 五、Spearman相关性分析
# ============================================================
def spearman_analysis(df):
    """分析特征与消费的Spearman相关性"""
    print("\n" + "=" * 80)
    print("正在进行Spearman相关性分析...")
    print("=" * 80)

    # 数值特征
    numeric_cols = df.select_dtypes(include=[np.number]).columns.tolist()
    target_features = ['累计消费(美元)', '是否付费', '是否高价值']

    spearman_results = []
    for target in target_features:
        for feature in numeric_cols:
            if feature == target or feature == '用户ID':
                continue
            # 过滤缺失值
            valid_data = df[[feature, target]].dropna()
            if len(valid_data) < 100:
                continue
            rho, p_value = spearmanr(valid_data[feature], valid_data[target])
            spearman_results.append({
                '目标变量': target,
                '特征': feature,
                'Spearman系数': rho,
                'P值': p_value,
                '显著': '是' if p_value < 0.05 else '否'
            })

    spearman_df = pd.DataFrame(spearman_results)

    # 分别输出每个目标变量的Top特征
    for target in target_features:
        target_df = spearman_df[spearman_df['目标变量'] == target]
        target_df = target_df.reindex(target_df['Spearman系数'].abs().sort_values(ascending=False).index)
        print(f"\n与 '{target}' 相关性Top 10 特征:")
        print(target_df.head(10).to_string(index=False))

    # 保存结果
    spearman_path = os.path.join(OUTPUT_DIR, '09_Spearman相关性分析结果.csv')
    spearman_df.to_csv(spearman_path, index=False, encoding='utf-8-sig')
    print(f"\nSpearman分析结果已保存: {spearman_path}")

    return spearman_df


# ============================================================
# 六、付费预测模型
# ============================================================
def prepare_modeling_data(df):
    """准备建模数据"""
    print("\n" + "=" * 80)
    print("正在准备付费预测模型数据...")
    print("=" * 80)

    target = '是否付费'

    # 建模特征：排除与目标变量有直接因果关系的数据泄露特征
    # 排除：VIP等级、送礼相关特征、累计消费等付费后/结果变量
    numeric_features = [
        '年龄', '关注主播数', '社交强度得分',
        '总观看时长(秒)_mean', '点赞次数_mean', '评论次数_mean',
        '关注次数_mean', '参与PK次数_mean',
        '连续登录天数_max', '是否跨时区观看_mean'
    ]

    categorical_features = [
        '性别', '设备类型', '网络类型', '地区', '用户类型', '是否主播'
    ]

    # 确保特征存在
    numeric_features = [f for f in numeric_features if f in df.columns]
    categorical_features = [f for f in categorical_features if f in df.columns]

    # 处理目标变量缺失
    df_model = df[df[target].notna()].copy()

    print(f"建模样本数: {len(df_model):,}")
    print(f"付费用户比例: {df_model[target].mean():.2%}")

    # 类别平衡采样
    if len(df_model) > 200000:
        paid = df_model[df_model[target] == 1]
        unpaid = df_model[df_model[target] == 0]
        sample_size = 100000
        paid_sample = paid.sample(n=min(sample_size, len(paid)), random_state=42)
        unpaid_sample = unpaid.sample(n=min(sample_size, len(unpaid)), random_state=42)
        df_model = pd.concat([paid_sample, unpaid_sample], ignore_index=True)
        print(f"采样后样本数: {len(df_model):,}")
        print(f"采样后付费用户比例: {df_model[target].mean():.2%}")

    X = df_model[numeric_features + categorical_features].copy()
    y = df_model[target].copy()

    return X, y, numeric_features, categorical_features


def build_preprocessor(numeric_features, categorical_features):
    """构建预处理管道"""
    numeric_transformer = Pipeline(steps=[
        ('imputer', SimpleImputer(strategy='median')),
        ('scaler', StandardScaler())
    ])

    categorical_transformer = Pipeline(steps=[
        ('imputer', SimpleImputer(strategy='most_frequent')),
        ('onehot', OneHotEncoder(handle_unknown='ignore', sparse_output=False))
    ])

    preprocessor = ColumnTransformer(
        transformers=[
            ('num', numeric_transformer, numeric_features),
            ('cat', categorical_transformer, categorical_features)
        ])

    return preprocessor


def evaluate_model(model, X_train, X_test, y_train, y_test, model_name):
    """评估模型"""
    print(f"\n{'=' * 80}")
    print(f"模型评估: {model_name}")
    print(f"{'=' * 80}")

    model.fit(X_train, y_train)

    y_pred_train = model.predict(X_train)
    y_pred_test = model.predict(X_test)

    if hasattr(model, "predict_proba"):
        y_prob_train = model.predict_proba(X_train)[:, 1]
        y_prob_test = model.predict_proba(X_test)[:, 1]
    elif hasattr(model, "decision_function"):
        y_prob_train = model.decision_function(X_train)
        y_prob_test = model.decision_function(X_test)
    else:
        y_prob_train = y_pred_train.astype(float)
        y_prob_test = y_pred_test.astype(float)

    metrics = {
        'model_name': model_name,
        'train_auc': roc_auc_score(y_train, y_prob_train),
        'test_auc': roc_auc_score(y_test, y_prob_test),
        'train_acc': accuracy_score(y_train, y_pred_train),
        'test_acc': accuracy_score(y_test, y_pred_test),
        'test_precision': precision_score(y_test, y_pred_test),
        'test_recall': recall_score(y_test, y_pred_test),
        'test_f1': f1_score(y_test, y_pred_test),
        'confusion_matrix': confusion_matrix(y_test, y_pred_test),
        'classification_report': classification_report(y_test, y_pred_test, target_names=['未付费', '付费']),
        'y_prob_test': y_prob_test,
        'y_test': y_test,
    }

    print(f"训练集 AUC: {metrics['train_auc']:.4f}")
    print(f"测试集 AUC: {metrics['test_auc']:.4f}")
    print(f"测试集 准确率: {metrics['test_acc']:.4f}")
    print(f"测试集 精确率: {metrics['test_precision']:.4f}")
    print(f"测试集 召回率: {metrics['test_recall']:.4f}")
    print(f"测试集 F1值: {metrics['test_f1']:.4f}")
    print(f"\n混淆矩阵:")
    print(metrics['confusion_matrix'])

    return model, metrics


def train_linearsvc(X_train, X_test, y_train, y_test, preprocessor):
    """训练LinearSVC模型"""
    print("\n" + "=" * 80)
    print("正在训练LinearSVC模型...")
    print("=" * 80)

    model = Pipeline(steps=[
        ('preprocessor', preprocessor),
        ('classifier', LinearSVC(random_state=42, max_iter=2000, dual='auto'))
    ])

    return evaluate_model(model, X_train, X_test, y_train, y_test, "LinearSVC")


def train_mlp(X_train, X_test, y_train, y_test, preprocessor):
    """训练MLP神经网络并调参"""
    print("\n" + "=" * 80)
    print("正在训练MLP神经网络（随机搜索调参）...")
    print("=" * 80)

    X_train_processed = preprocessor.fit_transform(X_train)
    X_test_processed = preprocessor.transform(X_test)

    param_dist = {
        'hidden_layer_sizes': [(64,), (128,), (128, 64)],
        'alpha': [1e-4, 1e-3, 1e-2],
        'learning_rate_init': [1e-3, 5e-3, 1e-2],
    }

    random_search = RandomizedSearchCV(
        MLPClassifier(max_iter=300, random_state=42, early_stopping=True),
        param_dist,
        n_iter=5,
        cv=3,
        scoring='roc_auc',
        n_jobs=1,
        random_state=42,
        verbose=1
    )

    random_search.fit(X_train_processed, y_train)
    print(f"\n最优参数: {random_search.best_params_}")
    print(f"最优交叉验证 AUC: {random_search.best_score_:.4f}")

    best_model = random_search.best_estimator_

    y_pred_train = best_model.predict(X_train_processed)
    y_pred_test = best_model.predict(X_test_processed)
    y_prob_train = best_model.predict_proba(X_train_processed)[:, 1]
    y_prob_test = best_model.predict_proba(X_test_processed)[:, 1]

    metrics = {
        'model_name': 'MLP神经网络',
        'best_params': str(random_search.best_params_),
        'cv_auc': random_search.best_score_,
        'train_auc': roc_auc_score(y_train, y_prob_train),
        'test_auc': roc_auc_score(y_test, y_prob_test),
        'train_acc': accuracy_score(y_train, y_pred_train),
        'test_acc': accuracy_score(y_test, y_pred_test),
        'test_precision': precision_score(y_test, y_pred_test),
        'test_recall': recall_score(y_test, y_pred_test),
        'test_f1': f1_score(y_test, y_pred_test),
        'confusion_matrix': confusion_matrix(y_test, y_pred_test),
        'classification_report': classification_report(y_test, y_pred_test, target_names=['未付费', '付费']),
        'y_prob_test': y_prob_test,
        'y_test': y_test,
    }

    print(f"训练集 AUC: {metrics['train_auc']:.4f}")
    print(f"测试集 AUC: {metrics['test_auc']:.4f}")
    print(f"测试集 准确率: {metrics['test_acc']:.4f}")
    print(f"\n混淆矩阵:")
    print(metrics['confusion_matrix'])

    return best_model, metrics


def get_feature_importance(model, preprocessor, numeric_features, categorical_features):
    """获取LinearSVC的特征重要性（系数绝对值）"""
    classifier = model.named_steps['classifier']
    feature_names = (numeric_features +
                     list(preprocessor.named_transformers_['cat']
                          .named_steps['onehot'].get_feature_names_out(categorical_features)))

    coefficients = classifier.coef_[0]
    importance_df = pd.DataFrame({
        '特征': feature_names,
        '系数': coefficients,
        '重要性': np.abs(coefficients)
    }).sort_values('重要性', ascending=False)

    return importance_df


# ============================================================
# 七、可视化
# ============================================================
def plot_confusion_matrix(metrics, save_path):
    """绘制混淆矩阵"""
    cm = metrics['confusion_matrix']
    fig, ax = plt.subplots(figsize=(8, 6))
    im = ax.imshow(cm, interpolation='nearest', cmap=plt.cm.Blues)
    ax.figure.colorbar(im, ax=ax)
    ax.set(xticks=[0, 1], yticks=[0, 1],
           xticklabels=['未付费', '付费'],
           yticklabels=['未付费', '付费'],
           title=f"{metrics['model_name']} 混淆矩阵",
           ylabel='真实标签', xlabel='预测标签')

    thresh = cm.max() / 2.
    for i in range(cm.shape[0]):
        for j in range(cm.shape[1]):
            ax.text(j, i, format(cm[i, j], 'd'),
                    ha="center", va="center",
                    color="white" if cm[i, j] > thresh else "black", fontsize=14)
    plt.tight_layout()
    plt.savefig(save_path, dpi=300, bbox_inches='tight')
    plt.close()


def plot_roc_curves(metrics_list, save_path):
    """绘制ROC曲线"""
    fig, ax = plt.subplots(figsize=(8, 6))
    for metrics in metrics_list:
        fpr, tpr, _ = roc_curve(metrics['y_test'], metrics['y_prob_test'])
        auc = metrics['test_auc']
        ax.plot(fpr, tpr, label=f"{metrics['model_name']} (AUC={auc:.4f})")
    ax.plot([0, 1], [0, 1], 'k--', label='随机猜测')
    ax.set_xlabel('假阳性率')
    ax.set_ylabel('真阳性率')
    ax.set_title('付费预测模型 ROC 曲线对比')
    ax.legend()
    ax.grid(True)
    plt.tight_layout()
    plt.savefig(save_path, dpi=300, bbox_inches='tight')
    plt.close()


def plot_feature_importance(importance_df, save_path, top_n=15):
    """绘制特征重要性"""
    top_df = importance_df.head(top_n)
    fig, ax = plt.subplots(figsize=(10, 8))
    colors = ['#2E86AB' if c >= 0 else '#A23B72' for c in top_df['系数']]
    ax.barh(range(len(top_df)), top_df['重要性'], color=colors)
    ax.set_yticks(range(len(top_df)))
    ax.set_yticklabels(top_df['特征'])
    ax.invert_yaxis()
    ax.set_xlabel('重要性（系数绝对值）')
    ax.set_title(f'LinearSVC Top {top_n} 特征重要性')
    plt.tight_layout()
    plt.savefig(save_path, dpi=300, bbox_inches='tight')
    plt.close()


def plot_spearman_heatmap(spearman_df, save_path):
    """绘制Spearman相关性热力图"""
    # 选择每个目标变量Top10特征
    top_features = set()
    for target in ['是否付费', '是否高价值', '累计消费(美元)']:
        target_df = spearman_df[spearman_df['目标变量'] == target]
        top = target_df.reindex(target_df['Spearman系数'].abs().sort_values(ascending=False).index).head(10)
        top_features.update(top['特征'].tolist())

    # 构建透视表
    pivot_data = []
    for feature in top_features:
        row = {'特征': feature}
        for target in ['是否付费', '是否高价值', '累计消费(美元)']:
            target_df = spearman_df[(spearman_df['目标变量'] == target) & (spearman_df['特征'] == feature)]
            if not target_df.empty:
                row[target] = target_df['Spearman系数'].values[0]
            else:
                row[target] = 0
        pivot_data.append(row)

    pivot_df = pd.DataFrame(pivot_data).set_index('特征')

    fig, ax = plt.subplots(figsize=(10, 12))
    sns.heatmap(pivot_df, annot=True, fmt='.2f', cmap='RdBu_r', center=0,
                linewidths=0.5, ax=ax)
    ax.set_title('Spearman 相关性热力图')
    plt.tight_layout()
    plt.savefig(save_path, dpi=300, bbox_inches='tight')
    plt.close()


# ============================================================
# 八、生成Word报告
# ============================================================
def create_word_report(metrics_list, cluster_profile, importance_df, spearman_df):
    """生成Word报告"""
    print("\n" + "=" * 80)
    print("正在生成Word分析报告...")
    print("=" * 80)

    doc = Document()
    title = doc.add_heading('问题2：高价值用户画像分析报告', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"报告生成时间: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph()

    # 一、分析目标
    doc.add_heading('一、分析目标', level=1)
    doc.add_paragraph(
        '本分析通过K-Means聚类、Spearman相关性分析和机器学习模型，'
        '识别高价值用户和付费用户的特征，为精准营销和用户运营提供依据。'
    )

    # 二、K-Means聚类结果
    doc.add_heading('二、K-Means聚类结果', level=1)
    doc.add_paragraph('不同用户群体的画像特征如下表所示：')

    table = doc.add_table(rows=1, cols=len(cluster_profile.columns))
    table.style = 'Light Grid Accent 1'
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(cluster_profile.columns):
        hdr_cells[i].text = str(h)

    for idx, row in cluster_profile.iterrows():
        row_cells = table.add_row().cells
        for i, val in enumerate(row):
            row_cells[i].text = f"{val:.4f}" if isinstance(val, float) else str(val)

    doc.add_paragraph()
    cluster_img = os.path.join(OUTPUT_DIR, '09_KMeans_聚类评估.png')
    if os.path.exists(cluster_img):
        doc.add_picture(cluster_img, width=Inches(5.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 三、付费预测模型结果
    doc.add_heading('三、付费预测模型结果', level=1)
    table2 = doc.add_table(rows=1, cols=7)
    table2.style = 'Light Grid Accent 1'
    hdr_cells2 = table2.rows[0].cells
    headers = ['模型', '测试集AUC', '测试集准确率', '精确率', '召回率', 'F1值', '训练集AUC']
    for i, h in enumerate(headers):
        hdr_cells2[i].text = h

    for metrics in metrics_list:
        row_cells = table2.add_row().cells
        row_cells[0].text = metrics['model_name']
        row_cells[1].text = f"{metrics['test_auc']:.4f}"
        row_cells[2].text = f"{metrics['test_acc']:.4f}"
        row_cells[3].text = f"{metrics['test_precision']:.4f}"
        row_cells[4].text = f"{metrics['test_recall']:.4f}"
        row_cells[5].text = f"{metrics['test_f1']:.4f}"
        row_cells[6].text = f"{metrics['train_auc']:.4f}"

    doc.add_paragraph()

    # 混淆矩阵
    doc.add_heading('四、混淆矩阵可视化', level=1)
    for metrics in metrics_list:
        doc.add_paragraph(f"{metrics['model_name']} 混淆矩阵:")
        img_path = os.path.join(OUTPUT_DIR, f'09_{metrics["model_name"]}_混淆矩阵.png')
        if os.path.exists(img_path):
            doc.add_picture(img_path, width=Inches(5.0))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ROC曲线
    doc.add_heading('五、ROC曲线对比', level=1)
    roc_path = os.path.join(OUTPUT_DIR, '09_ROC曲线对比.png')
    if os.path.exists(roc_path):
        doc.add_picture(roc_path, width=Inches(5.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 六、特征重要性
    doc.add_heading('六、LinearSVC特征重要性', level=1)
    doc.add_paragraph('基于LinearSVC模型系数绝对值排序的Top 15重要特征：')

    top15 = importance_df.head(15)
    table3 = doc.add_table(rows=1, cols=3)
    table3.style = 'Light Grid Accent 1'
    hdr_cells3 = table3.rows[0].cells
    for i, h in enumerate(['排名', '特征', '重要性']):
        hdr_cells3[i].text = h

    for idx, (_, row) in enumerate(top15.iterrows(), 1):
        row_cells = table3.add_row().cells
        row_cells[0].text = str(idx)
        row_cells[1].text = str(row['特征'])
        row_cells[2].text = f"{row['重要性']:.4f}"

    doc.add_paragraph()
    importance_img = os.path.join(OUTPUT_DIR, '09_LinearSVC_特征重要性.png')
    if os.path.exists(importance_img):
        doc.add_picture(importance_img, width=Inches(5.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 七、Spearman相关性
    doc.add_heading('七、Spearman相关性分析', level=1)
    spearman_img = os.path.join(OUTPUT_DIR, '09_Spearman相关性热力图.png')
    if os.path.exists(spearman_img):
        doc.add_picture(spearman_img, width=Inches(5.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 八、结论
    doc.add_heading('八、结论与建议', level=1)
    best_model = max(metrics_list, key=lambda x: x['test_auc'])
    doc.add_paragraph(f"1. 最优付费预测模型为 {best_model['model_name']}，测试集 AUC 为 {best_model['test_auc']:.4f}。")

    if not importance_df.empty:
        top1 = importance_df.iloc[0]
        doc.add_paragraph(f"2. 最重要的付费驱动特征: {top1['特征']} (重要性: {top1['重要性']:.4f})。")

    doc.add_paragraph(
        '3. 建议运营侧重点关注高互动、高观看时长、高VIP等级的用户群体，'
        '并通过聚类结果制定差异化的用户运营策略。'
    )

    doc.save(DOC_FILE)
    print(f"Word报告已保存: {DOC_FILE}")


# ============================================================
# 九、主程序
# ============================================================
def main():
    sys.stdout = Tee(LOG_FILE)

    print("\n" + "=" * 80)
    print("问题2：高价值用户画像")
    print("=" * 80)
    print(f"开始时间: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    # 1. 加载与特征工程
    df = load_and_engineer_features()

    # 2. K-Means聚类
    df, cluster_profile = kmeans_clustering(df)

    # 3. Spearman相关性分析
    spearman_df = spearman_analysis(df)

    # 4. 准备付费预测数据
    X, y, numeric_features, categorical_features = prepare_modeling_data(df)
    X_train, X_test, y_train, y_test = train_test_split(
        X, y, test_size=0.2, random_state=42, stratify=y
    )
    print(f"\n训练集样本数: {len(X_train):,}")
    print(f"测试集样本数: {len(X_test):,}")

    preprocessor = build_preprocessor(numeric_features, categorical_features)

    # 5. LinearSVC
    svm_model, svm_metrics = train_linearsvc(
        X_train, X_test, y_train, y_test, preprocessor
    )

    # 6. MLP
    mlp_model, mlp_metrics = train_mlp(
        X_train, X_test, y_train, y_test, preprocessor
    )

    # 7. 特征重要性
    importance_df = get_feature_importance(svm_model, preprocessor,
                                           numeric_features, categorical_features)
    print("\nLinearSVC Top 10 特征重要性:")
    print(importance_df.head(10).to_string(index=False))

    # 8. 可视化
    metrics_list = [svm_metrics, mlp_metrics]
    plot_confusion_matrix(svm_metrics, os.path.join(OUTPUT_DIR, '09_LinearSVC_混淆矩阵.png'))
    plot_confusion_matrix(mlp_metrics, os.path.join(OUTPUT_DIR, '09_MLP神经网络_混淆矩阵.png'))
    plot_roc_curves(metrics_list, os.path.join(OUTPUT_DIR, '09_ROC曲线对比.png'))
    plot_feature_importance(importance_df, os.path.join(OUTPUT_DIR, '09_LinearSVC_特征重要性.png'))
    plot_spearman_heatmap(spearman_df, os.path.join(OUTPUT_DIR, '09_Spearman相关性热力图.png'))

    # 9. Word报告
    create_word_report(metrics_list, cluster_profile, importance_df, spearman_df)

    print("\n" + "=" * 80)
    print("问题2分析完成！")
    print(f"结束时间: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print("=" * 80)


if __name__ == '__main__':
    main()
