# -*- coding: utf-8 -*-
"""
问题4：AB实验效果评估
方法：分组均衡性检查 + Wilcoxon秩和检验 + 卡方检验
目标：评估新用户引导实验对留存率等指标的影响
"""

import os
import sys
import warnings
import datetime
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from scipy.stats import chi2_contingency, mannwhitneyu
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

warnings.filterwarnings('ignore')

# ============================================================
# 一、配置
# ============================================================
_CODE_DIR = os.path.dirname(os.path.abspath(__file__))
_ROOT_DIR = os.path.dirname(_CODE_DIR)
DATA_DIR = os.path.join(_ROOT_DIR, 'data', '中文版')
OUTPUT_DIR = os.path.join(_CODE_DIR, 'output')
DOC_DIR = os.path.join(_ROOT_DIR, 'reports')
LOG_FILE = os.path.join(OUTPUT_DIR, '11_问题4_分析日志.txt')
DOC_FILE = os.path.join(DOC_DIR, '11_问题4_AB实验分析报告.docx')

os.makedirs(OUTPUT_DIR, exist_ok=True)

plt.rcParams['font.sans-serif'] = ['SimHei', 'Microsoft YaHei']
plt.rcParams['axes.unicode_minus'] = False


# ============================================================
# 二、日志工具
# ============================================================
class Tee:
    def __init__(self, filename):
        self.terminal = sys.stdout
        self.log = open(filename, 'w', encoding='utf-8')

    def write(self, message):
        self.terminal.write(message)
        self.log.write(message)
        self.flush()

    def flush(self):
        self.terminal.flush()
        self.log.flush()


# ============================================================
# 三、数据加载
# ============================================================
def load_ab_data():
    """加载AB实验数据"""
    print("=" * 80)
    print("正在加载AB实验数据...")
    print("=" * 80)

    ab_path = os.path.join(DATA_DIR, 'AB实验曝光表.csv')
    df = pd.read_csv(ab_path)
    print(f"AB实验曝光表: {df.shape[0]:,} 行 × {df.shape[1]} 列")
    print(f"\n字段列表: {list(df.columns)}")

    # 查看分组情况
    print(f"\n实验分组:")
    print(df['实验分组'].value_counts())

    return df


# ============================================================
# 四、分组均衡性检查
# ============================================================
def balance_check(df):
    """检查实验分组和对照组在关键特征上的均衡性"""
    print("\n" + "=" * 80)
    print("正在进行分组均衡性检查...")
    print("=" * 80)

    treatment = df[df['实验分组'] == 'treatment']
    control = df[df['实验分组'] == 'control']

    print(f"实验分组样本数: {len(treatment):,}")
    print(f"对照组样本数: {len(control):,}")

    # 检查分组变量
    balance_features = []
    for col in df.columns:
        if col in ['用户ID', '实验分组']:
            continue
        # 数值型：比较均值
        if df[col].dtype in ['int64', 'float64']:
            t_mean = treatment[col].mean()
            c_mean = control[col].mean()
            diff_pct = (t_mean - c_mean) / c_mean if c_mean != 0 else 0
            balance_features.append({
                '特征': col,
                '类型': '数值',
                '实验分组均值': t_mean,
                '对照组均值': c_mean,
                '差异百分比': diff_pct,
            })
        # 分类型：比较分布
        else:
            t_dist = treatment[col].value_counts(normalize=True)
            c_dist = control[col].value_counts(normalize=True)
            balance_features.append({
                '特征': col,
                '类型': '分类',
                '实验分组分布': str(dict(t_dist.head(3))),
                '对照组分布': str(dict(c_dist.head(3))),
                '差异百分比': np.nan,
            })

    balance_df = pd.DataFrame(balance_features)
    print("\n分组均衡性:")
    print(balance_df.to_string(index=False))

    # 保存
    balance_path = os.path.join(OUTPUT_DIR, '11_分组均衡性检查.csv')
    balance_df.to_csv(balance_path, index=False, encoding='utf-8-sig')
    print(f"\n均衡性结果已保存: {balance_path}")

    return balance_df


# ============================================================
# 五、Wilcoxon秩和检验
# ============================================================
def wilcoxon_test(df):
    """对连续型指标进行Wilcoxon秩和检验"""
    print("\n" + "=" * 80)
    print("正在进行Wilcoxon秩和检验...")
    print("=" * 80)

    treatment = df[df['实验分组'] == 'treatment']
    control = df[df['实验分组'] == 'control']

    # 自动识别数值型指标（排除ID和分组）
    metric_cols = [col for col in df.columns
                   if col not in ['用户ID', '实验分组']
                   and df[col].dtype in ['int64', 'float64']]

    wilcoxon_results = []

    for metric in metric_cols:
        t_values = treatment[metric].dropna()
        c_values = control[metric].dropna()

        if len(t_values) < 10 or len(c_values) < 10:
            continue

        # Wilcoxon秩和检验（Mann-Whitney U）
        statistic, p_value = mannwhitneyu(
            t_values, c_values, alternative='two-sided'
        )

        t_median = t_values.median()
        c_median = c_values.median()
        lift = (t_median - c_median) / c_median if c_median != 0 else 0

        wilcoxon_results.append({
            '指标': metric,
            '实验分组中位数': t_median,
            '对照组中位数': c_median,
            '绝对差异': t_median - c_median,
            '相对提升': lift,
            '统计量': statistic,
            'P值': p_value,
            '显著': '是' if p_value < 0.05 else '否'
        })

    wilcoxon_df = pd.DataFrame(wilcoxon_results)
    if not wilcoxon_df.empty:
        wilcoxon_df = wilcoxon_df.sort_values('P值')
        print("\nWilcoxon秩和检验结果:")
        print(wilcoxon_df.to_string(index=False))

        # 保存
        wilcoxon_path = os.path.join(OUTPUT_DIR, '11_Wilcoxon秩和检验结果.csv')
        wilcoxon_df.to_csv(wilcoxon_path, index=False, encoding='utf-8-sig')
        print(f"\nWilcoxon结果已保存: {wilcoxon_path}")
    else:
        print("未找到合适的连续型指标进行Wilcoxon检验。")

    return wilcoxon_df


# ============================================================
# 六、卡方检验
# ============================================================
def chi_square_test(df):
    """对二分类指标进行卡方检验"""
    print("\n" + "=" * 80)
    print("正在进行卡方检验...")
    print("=" * 80)

    treatment = df[df['实验分组'] == 'treatment']
    control = df[df['实验分组'] == 'control']

    # 自动识别可能的二分类指标（0/1或True/False）
    binary_cols = []
    for col in df.columns:
        if col in ['用户ID', '实验分组']:
            continue
        unique_vals = df[col].dropna().unique()
        if set(unique_vals).issubset({0, 1, True, False}):
            binary_cols.append(col)

    chi2_results = []

    for metric in binary_cols:
        contingency = pd.crosstab(df['实验分组'], df[metric])
        if contingency.shape != (2, 2):
            continue

        chi2, p_value, dof, expected = chi2_contingency(contingency)

        t_rate = treatment[metric].mean()
        c_rate = control[metric].mean()
        lift = (t_rate - c_rate) / c_rate if c_rate != 0 else 0

        chi2_results.append({
            '指标': metric,
            '实验分组比例': t_rate,
            '对照组比例': c_rate,
            '绝对差异': t_rate - c_rate,
            '相对提升': lift,
            '卡方值': chi2,
            'P值': p_value,
            '显著': '是' if p_value < 0.05 else '否'
        })

    chi2_df = pd.DataFrame(chi2_results)
    if not chi2_df.empty:
        chi2_df = chi2_df.sort_values('P值')
        print("\n卡方检验结果:")
        print(chi2_df.to_string(index=False))

        # 保存
        chi2_path = os.path.join(OUTPUT_DIR, '11_卡方检验结果.csv')
        chi2_df.to_csv(chi2_path, index=False, encoding='utf-8-sig')
        print(f"\n卡方检验结果已保存: {chi2_path}")
    else:
        print("未找到合适的二分类指标进行卡方检验。")

    return chi2_df


# ============================================================
# 七、可视化
# ============================================================
def plot_metric_comparison(df, metrics, save_path):
    """绘制关键指标的组间对比图"""
    if not metrics:
        return

    n_metrics = len(metrics)
    fig, axes = plt.subplots(1, n_metrics, figsize=(6 * n_metrics, 5))
    if n_metrics == 1:
        axes = [axes]

    for ax, metric in zip(axes, metrics):
        data_to_plot = [df[df['实验分组'] == 'treatment'][metric].dropna(),
                        df[df['实验分组'] == 'control'][metric].dropna()]
        ax.boxplot(data_to_plot, labels=['实验分组', '对照组'])
        ax.set_ylabel(metric)
        ax.set_title(f'{metric} 组间对比')
        ax.grid(True, alpha=0.3)

    plt.tight_layout()
    plt.savefig(save_path, dpi=300, bbox_inches='tight')
    plt.close()


def plot_retention_rates(df, metrics, save_path):
    """绘制留存率对比柱状图"""
    if not metrics:
        return

    treatment = df[df['实验分组'] == 'treatment']
    control = df[df['实验分组'] == 'control']

    t_rates = [treatment[m].mean() for m in metrics]
    c_rates = [control[m].mean() for m in metrics]

    x = np.arange(len(metrics))
    width = 0.35

    fig, ax = plt.subplots(figsize=(10, 6))
    ax.bar(x - width/2, t_rates, width, label='实验分组', color='#2E86AB')
    ax.bar(x + width/2, c_rates, width, label='对照组', color='#A23B72')

    ax.set_ylabel('比例')
    ax.set_title('实验分组 vs 对照组 留存率对比')
    ax.set_xticks(x)
    ax.set_xticklabels(metrics)
    ax.legend()
    ax.grid(True, axis='y', alpha=0.3)

    # 添加数值标签
    for i, (t, c) in enumerate(zip(t_rates, c_rates)):
        ax.text(i - width/2, t, f'{t:.2%}', ha='center', va='bottom', fontsize=9)
        ax.text(i + width/2, c, f'{c:.2%}', ha='center', va='bottom', fontsize=9)

    plt.tight_layout()
    plt.savefig(save_path, dpi=300, bbox_inches='tight')
    plt.close()


# ============================================================
# 八、生成Word报告
# ============================================================
def create_word_report(balance_df, wilcoxon_df, chi2_df, df):
    """生成Word报告"""
    print("\n" + "=" * 80)
    print("正在生成Word分析报告...")
    print("=" * 80)

    doc = Document()
    title = doc.add_heading('问题4：AB实验效果评估报告', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"报告生成时间: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph()

    # 一、分析目标
    doc.add_heading('一、分析目标', level=1)
    doc.add_paragraph(
        '本分析通过Wilcoxon秩和检验和卡方检验，评估新用户引导实验（treatment vs control）'
        '对核心留存指标和连续指标的影响。'
    )

    # 二、样本概览
    doc.add_heading('二、样本概览', level=1)
    treatment = df[df['实验分组'] == 'treatment']
    control = df[df['实验分组'] == 'control']
    doc.add_paragraph(f"实验分组样本数: {len(treatment):,}")
    doc.add_paragraph(f"对照组样本数: {len(control):,}")

    # 三、分组均衡性
    doc.add_heading('三、分组均衡性检查', level=1)
    if not balance_df.empty:
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Light Grid Accent 1'
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(['特征', '类型', '实验分组均值/分布', '对照组均值/分布', '差异百分比']):
            hdr_cells[i].text = h

        for _, row in balance_df.iterrows():
            row_cells = table.add_row().cells
            row_cells[0].text = str(row['特征'])
            row_cells[1].text = str(row['类型'])
            if row['类型'] == '数值':
                row_cells[2].text = f"{row['实验分组均值']:.4f}"
                row_cells[3].text = f"{row['对照组均值']:.4f}"
            else:
                row_cells[2].text = str(row['实验分组分布'])[:50]
                row_cells[3].text = str(row['对照组分布'])[:50]
            row_cells[4].text = f"{row['差异百分比']:.2%}" if pd.notna(row['差异百分比']) else '-'

    # 四、Wilcoxon检验结果
    if not wilcoxon_df.empty:
        doc.add_heading('四、连续指标 Wilcoxon 秩和检验', level=1)
        table2 = doc.add_table(rows=1, cols=6)
        table2.style = 'Light Grid Accent 1'
        hdr_cells2 = table2.rows[0].cells
        for i, h in enumerate(['指标', '实验分组中位数', '对照组中位数', '相对提升', 'P值', '显著']):
            hdr_cells2[i].text = h

        for _, row in wilcoxon_df.iterrows():
            row_cells = table2.add_row().cells
            row_cells[0].text = str(row['指标'])
            row_cells[1].text = f"{row['实验分组中位数']:.4f}"
            row_cells[2].text = f"{row['对照组中位数']:.4f}"
            row_cells[3].text = f"{row['相对提升']:.2%}"
            row_cells[4].text = f"{row['P值']:.4f}"
            row_cells[5].text = str(row['显著'])

        doc.add_paragraph()
        metric_img = os.path.join(OUTPUT_DIR, '11_连续指标对比.png')
        if os.path.exists(metric_img):
            doc.add_picture(metric_img, width=Inches(5.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 五、卡方检验结果
    if not chi2_df.empty:
        doc.add_heading('五、二分类指标 卡方检验', level=1)
        table3 = doc.add_table(rows=1, cols=6)
        table3.style = 'Light Grid Accent 1'
        hdr_cells3 = table3.rows[0].cells
        for i, h in enumerate(['指标', '实验分组比例', '对照组比例', '相对提升', 'P值', '显著']):
            hdr_cells3[i].text = h

        for _, row in chi2_df.iterrows():
            row_cells = table3.add_row().cells
            row_cells[0].text = str(row['指标'])
            row_cells[1].text = f"{row['实验分组比例']:.2%}"
            row_cells[2].text = f"{row['对照组比例']:.2%}"
            row_cells[3].text = f"{row['相对提升']:.2%}"
            row_cells[4].text = f"{row['P值']:.4f}"
            row_cells[5].text = str(row['显著'])

        doc.add_paragraph()
        retention_img = os.path.join(OUTPUT_DIR, '11_留存率对比.png')
        if os.path.exists(retention_img):
            doc.add_picture(retention_img, width=Inches(5.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 六、结论
    doc.add_heading('六、结论与建议', level=1)

    if not wilcoxon_df.empty and wilcoxon_df['显著'].str.contains('是').any():
        sig_wilcoxon = wilcoxon_df[wilcoxon_df['显著'] == '是'].iloc[0]
        doc.add_paragraph(
            f"1. 连续指标中，'{sig_wilcoxon['指标']}' 在实验分组和对照组间存在显著差异 "
            f"(P={sig_wilcoxon['P值']:.4f})，相对提升 {sig_wilcoxon['相对提升']:.2%}。"
        )
    else:
        doc.add_paragraph("1. 连续指标中未发现实验分组与对照组存在显著差异。")

    if not chi2_df.empty and chi2_df['显著'].str.contains('是').any():
        sig_chi2 = chi2_df[chi2_df['显著'] == '是'].iloc[0]
        doc.add_paragraph(
            f"2. 二分类指标中，'{sig_chi2['指标']}' 在实验分组和对照组间存在显著差异 "
            f"(P={sig_chi2['P值']:.4f})，相对提升 {sig_chi2['相对提升']:.2%}。"
        )
    else:
        doc.add_paragraph("2. 二分类指标中未发现实验分组与对照组存在显著差异。")

    doc.add_paragraph(
        '3. 建议根据实验结论决定是否全量上线新用户引导策略，并持续监控长期留存效果。'
    )

    doc.save(DOC_FILE)
    print(f"Word报告已保存: {DOC_FILE}")


# ============================================================
# 九、主程序
# ============================================================
def main():
    sys.stdout = Tee(LOG_FILE)

    print("\n" + "=" * 80)
    print("问题4：AB实验效果评估")
    print("=" * 80)
    print(f"开始时间: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    # 1. 加载数据
    df = load_ab_data()

    # 2. 分组均衡性检查
    balance_df = balance_check(df)

    # 3. Wilcoxon秩和检验
    wilcoxon_df = wilcoxon_test(df)

    # 4. 卡方检验
    chi2_df = chi_square_test(df)

    # 5. 可视化
    if not wilcoxon_df.empty:
        top_metrics = wilcoxon_df.head(3)['指标'].tolist()
        plot_metric_comparison(df, top_metrics,
                               os.path.join(OUTPUT_DIR, '11_连续指标对比.png'))

    if not chi2_df.empty:
        retention_metrics = chi2_df['指标'].tolist()
        plot_retention_rates(df, retention_metrics,
                             os.path.join(OUTPUT_DIR, '11_留存率对比.png'))

    # 6. Word报告
    create_word_report(balance_df, wilcoxon_df, chi2_df, df)

    print("\n" + "=" * 80)
    print("问题4分析完成！")
    print(f"结束时间: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print("=" * 80)


if __name__ == '__main__':
    main()
