# -*- coding: utf-8 -*-
"""
问题3：直播间转化分析
方法：线性回归 + MLP神经网络回归 + 特征重要性 + 残差分析
目标：识别高收入直播间的特征，预测直播间收入
"""

import os
import sys
import warnings
import datetime
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from sklearn.model_selection import train_test_split, RandomizedSearchCV
from sklearn.preprocessing import StandardScaler, OneHotEncoder
from sklearn.compose import ColumnTransformer
from sklearn.pipeline import Pipeline
from sklearn.linear_model import LinearRegression, Ridge
from sklearn.neural_network import MLPRegressor
from sklearn.metrics import mean_squared_error, mean_absolute_error, r2_score
from sklearn.impute import SimpleImputer
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

warnings.filterwarnings('ignore')

# ============================================================
# 一、配置
# ============================================================
_CODE_DIR = os.path.dirname(os.path.abspath(__file__))
_ROOT_DIR = os.path.dirname(_CODE_DIR)
DATA_DIR = os.path.join(_ROOT_DIR, 'data', '中文版')
OUTPUT_DIR = os.path.join(_CODE_DIR, 'output')
DOC_DIR = os.path.join(_ROOT_DIR, 'reports')
LOG_FILE = os.path.join(OUTPUT_DIR, '10_问题3_分析日志.txt')
DOC_FILE = os.path.join(DOC_DIR, '10_问题3_直播间转化分析报告.docx')

os.makedirs(OUTPUT_DIR, exist_ok=True)
os.environ['JOBLIB_TEMP_FOLDER'] = os.path.abspath('./joblib_temp')
os.makedirs(os.environ['JOBLIB_TEMP_FOLDER'], exist_ok=True)
os.environ['LOKY_MAX_CPU_COUNT'] = '4'
os.environ['LOKY_PICKLER'] = 'pickle'

plt.rcParams['font.sans-serif'] = ['SimHei', 'Microsoft YaHei']
plt.rcParams['axes.unicode_minus'] = False


# ============================================================
# 二、日志工具
# ============================================================
class Tee:
    def __init__(self, filename):
        self.terminal = sys.stdout
        self.log = open(filename, 'w', encoding='utf-8')

    def write(self, message):
        self.terminal.write(message)
        self.log.write(message)
        self.flush()

    def flush(self):
        self.terminal.flush()
        self.log.flush()


# ============================================================
# 三、数据加载与特征工程
# ============================================================
def load_and_engineer_features():
    """加载直播间数据并关联主播特征"""
    print("=" * 80)
    print("正在加载直播间数据并构建特征...")
    print("=" * 80)

    # 1. 直播间日统计表
    room_path = os.path.join(DATA_DIR, '直播间日统计.csv')
    df_room = pd.read_csv(room_path)
    print(f"直播间日统计表: {df_room.shape[0]:,} 行 × {df_room.shape[1]} 列")

    # 2. 主播基础信息表
    streamer_path = os.path.join(DATA_DIR, '主播基础信息.csv')
    df_streamer = pd.read_csv(streamer_path)
    print(f"主播基础信息表: {df_streamer.shape[0]:,} 行 × {df_streamer.shape[1]} 列")

    # 合并直播间与主播信息
    df = df_room.merge(df_streamer, on='主播ID', how='left')

    # 3. 特征工程
    # 互动转化率（直播间日统计表无送礼人数，用互动总数替代）
    df['互动转化率'] = np.where(df['观看人数'] > 0,
                                df['互动总数'] / df['观看人数'], 0)
    # 直播时长（秒）
    df['直播时长(秒)'] = df['直播时长(小时)'] * 3600
    # 是否高收入（用于分类参考）
    df['是否高收入'] = (df['礼物收入(美元)'] >= df['礼物收入(美元)'].quantile(0.8)).astype(int)

    print(f"\n合并后数据: {df.shape[0]:,} 行 × {df.shape[1]} 列")
    print(f"平均礼物收入: ${df['礼物收入(美元)'].mean():.4f}")
    print(f"高收入直播间比例: {df['是否高收入'].mean():.2%}")

    return df


# ============================================================
# 四、准备建模数据
# ============================================================
def prepare_modeling_data(df):
    """准备回归建模数据"""
    print("\n" + "=" * 80)
    print("正在准备直播间收入预测模型数据...")
    print("=" * 80)

    target = '礼物收入(美元)'

    # 建模特征：排除与目标变量有直接因果关系的数据泄露特征
    # 排除：累计收到礼物(美元)、近两周收入变化率 等结果变量
    numeric_features = [
        '直播时长(秒)', '观看人数', '互动总数',
        '互动密度', '互动转化率', '跨地区观众占比',
        '粉丝数', 'PK胜率', '观众留存率',
        '日均直播时长', '是否新主播', '首周开播天数', '观看高峰小时'
    ]

    categorical_features = [
        '直播内容分类', '地区', '是否开启PK', '主播等级'
    ]

    # 确保特征存在
    numeric_features = [f for f in numeric_features if f in df.columns]
    categorical_features = [f for f in categorical_features if f in df.columns]

    print(f"数值特征 ({len(numeric_features)}): {numeric_features}")
    print(f"分类特征 ({len(categorical_features)}): {categorical_features}")

    # 过滤收入异常值和缺失目标
    df_model = df[df[target].notna() & (df[target] >= 0)].copy()

    print(f"建模样本数: {len(df_model):,}")
    print(f"平均礼物收入: ${df_model[target].mean():.4f}")

    # 数据量较大时采样
    if len(df_model) > 100000:
        df_model = df_model.sample(n=100000, random_state=42)
        print(f"采样后样本数: {len(df_model):,}")

    X = df_model[numeric_features + categorical_features].copy()
    y = df_model[target].copy()

    return X, y, numeric_features, categorical_features, df_model


def build_preprocessor(numeric_features, categorical_features):
    """构建预处理管道"""
    numeric_transformer = Pipeline(steps=[
        ('imputer', SimpleImputer(strategy='median')),
        ('scaler', StandardScaler())
    ])

    categorical_transformer = Pipeline(steps=[
        ('imputer', SimpleImputer(strategy='most_frequent')),
        ('onehot', OneHotEncoder(handle_unknown='ignore', sparse_output=False))
    ])

    preprocessor = ColumnTransformer(
        transformers=[
            ('num', numeric_transformer, numeric_features),
            ('cat', categorical_transformer, categorical_features)
        ])

    return preprocessor


# ============================================================
# 五、模型训练与评估
# ============================================================
def evaluate_regression_model(model, X_train, X_test, y_train, y_test, model_name):
    """评估回归模型"""
    print(f"\n{'=' * 80}")
    print(f"模型评估: {model_name}")
    print(f"{'=' * 80}")

    model.fit(X_train, y_train)

    y_pred_train = model.predict(X_train)
    y_pred_test = model.predict(X_test)

    # 防止负值预测
    y_pred_train = np.maximum(y_pred_train, 0)
    y_pred_test = np.maximum(y_pred_test, 0)

    metrics = {
        'model_name': model_name,
        'train_rmse': np.sqrt(mean_squared_error(y_train, y_pred_train)),
        'test_rmse': np.sqrt(mean_squared_error(y_test, y_pred_test)),
        'train_mae': mean_absolute_error(y_train, y_pred_train),
        'test_mae': mean_absolute_error(y_test, y_pred_test),
        'train_r2': r2_score(y_train, y_pred_train),
        'test_r2': r2_score(y_test, y_pred_test),
        'y_test': y_test,
        'y_pred_test': y_pred_test,
    }

    print(f"训练集 RMSE: {metrics['train_rmse']:.4f}")
    print(f"测试集 RMSE: {metrics['test_rmse']:.4f}")
    print(f"训练集 MAE: {metrics['train_mae']:.4f}")
    print(f"测试集 MAE: {metrics['test_mae']:.4f}")
    print(f"训练集 R²: {metrics['train_r2']:.4f}")
    print(f"测试集 R²: {metrics['test_r2']:.4f}")

    return model, metrics


def train_linear_regression(X_train, X_test, y_train, y_test, preprocessor):
    """训练线性回归模型"""
    print("\n" + "=" * 80)
    print("正在训练线性回归模型...")
    print("=" * 80)

    model = Pipeline(steps=[
        ('preprocessor', preprocessor),
        ('regressor', LinearRegression())
    ])

    return evaluate_regression_model(model, X_train, X_test, y_train, y_test, "线性回归")


def train_ridge(X_train, X_test, y_train, y_test, preprocessor):
    """训练Ridge回归模型"""
    print("\n" + "=" * 80)
    print("正在训练Ridge回归模型...")
    print("=" * 80)

    model = Pipeline(steps=[
        ('preprocessor', preprocessor),
        ('regressor', Ridge(alpha=1.0, random_state=42))
    ])

    return evaluate_regression_model(model, X_train, X_test, y_train, y_test, "Ridge回归")


def train_mlp_regression(X_train, X_test, y_train, y_test, preprocessor):
    """训练MLP神经网络回归并调参"""
    print("\n" + "=" * 80)
    print("正在训练MLP神经网络回归（随机搜索调参）...")
    print("=" * 80)

    X_train_processed = preprocessor.fit_transform(X_train)
    X_test_processed = preprocessor.transform(X_test)

    param_dist = {
        'hidden_layer_sizes': [(64,), (128,), (128, 64)],
        'alpha': [1e-4, 1e-3, 1e-2],
        'learning_rate_init': [1e-3, 5e-3, 1e-2],
    }

    random_search = RandomizedSearchCV(
        MLPRegressor(max_iter=300, random_state=42, early_stopping=True),
        param_dist,
        n_iter=5,
        cv=3,
        scoring='neg_mean_squared_error',
        n_jobs=1,
        random_state=42,
        verbose=1
    )

    random_search.fit(X_train_processed, y_train)
    print(f"\n最优参数: {random_search.best_params_}")
    print(f"最优交叉验证 RMSE: {np.sqrt(-random_search.best_score_):.4f}")

    best_model = random_search.best_estimator_

    y_pred_train = best_model.predict(X_train_processed)
    y_pred_test = best_model.predict(X_test_processed)
    y_pred_train = np.maximum(y_pred_train, 0)
    y_pred_test = np.maximum(y_pred_test, 0)

    metrics = {
        'model_name': 'MLP神经网络回归',
        'best_params': str(random_search.best_params_),
        'cv_rmse': np.sqrt(-random_search.best_score_),
        'train_rmse': np.sqrt(mean_squared_error(y_train, y_pred_train)),
        'test_rmse': np.sqrt(mean_squared_error(y_test, y_pred_test)),
        'train_mae': mean_absolute_error(y_train, y_pred_train),
        'test_mae': mean_absolute_error(y_test, y_pred_test),
        'train_r2': r2_score(y_train, y_pred_train),
        'test_r2': r2_score(y_test, y_pred_test),
        'y_test': y_test,
        'y_pred_test': y_pred_test,
    }

    print(f"训练集 RMSE: {metrics['train_rmse']:.4f}")
    print(f"测试集 RMSE: {metrics['test_rmse']:.4f}")
    print(f"训练集 R²: {metrics['train_r2']:.4f}")
    print(f"测试集 R²: {metrics['test_r2']:.4f}")

    return best_model, metrics


def get_feature_importance(model, preprocessor, numeric_features, categorical_features):
    """获取线性回归的特征重要性（系数绝对值）"""
    regressor = model.named_steps['regressor']
    feature_names = (numeric_features +
                     list(preprocessor.named_transformers_['cat']
                          .named_steps['onehot'].get_feature_names_out(categorical_features)))

    coefficients = regressor.coef_
    importance_df = pd.DataFrame({
        '特征': feature_names,
        '系数': coefficients,
        '重要性': np.abs(coefficients)
    }).sort_values('重要性', ascending=False)

    return importance_df


# ============================================================
# 六、可视化
# ============================================================
def plot_feature_importance(importance_df, save_path, top_n=15):
    """绘制特征重要性"""
    top_df = importance_df.head(top_n)
    fig, ax = plt.subplots(figsize=(10, 8))
    colors = ['#2E86AB' if c >= 0 else '#A23B72' for c in top_df['系数']]
    ax.barh(range(len(top_df)), top_df['重要性'], color=colors)
    ax.set_yticks(range(len(top_df)))
    ax.set_yticklabels(top_df['特征'])
    ax.invert_yaxis()
    ax.set_xlabel('重要性（系数绝对值）')
    ax.set_title(f'线性回归 Top {top_n} 特征重要性')
    plt.tight_layout()
    plt.savefig(save_path, dpi=300, bbox_inches='tight')
    plt.close()


def plot_actual_vs_predicted(metrics_list, save_path):
    """绘制真实值vs预测值散点图"""
    fig, axes = plt.subplots(1, len(metrics_list), figsize=(6 * len(metrics_list), 5))
    if len(metrics_list) == 1:
        axes = [axes]

    for ax, metrics in zip(axes, metrics_list):
        y_test = metrics['y_test']
        y_pred = metrics['y_pred_test']
        ax.scatter(y_test, y_pred, alpha=0.3, s=10)
        max_val = max(y_test.max(), y_pred.max())
        ax.plot([0, max_val], [0, max_val], 'r--', label='理想预测')
        ax.set_xlabel('真实礼物收入 ($)')
        ax.set_ylabel('预测礼物收入 ($)')
        ax.set_title(f"{metrics['model_name']}\nR²={metrics['test_r2']:.4f}")
        ax.legend()
        ax.grid(True, alpha=0.3)

    plt.tight_layout()
    plt.savefig(save_path, dpi=300, bbox_inches='tight')
    plt.close()


def plot_residuals(metrics, save_path):
    """绘制残差图"""
    residuals = metrics['y_test'] - metrics['y_pred_test']
    fig, ax = plt.subplots(figsize=(8, 6))
    ax.scatter(metrics['y_pred_test'], residuals, alpha=0.3, s=10)
    ax.axhline(y=0, color='r', linestyle='--')
    ax.set_xlabel('预测礼物收入 ($)')
    ax.set_ylabel('残差 ($)')
    ax.set_title(f"{metrics['model_name']} 残差图")
    ax.grid(True, alpha=0.3)
    plt.tight_layout()
    plt.savefig(save_path, dpi=300, bbox_inches='tight')
    plt.close()


def plot_income_distribution(df, save_path):
    """绘制礼物收入分布"""
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(12, 5))

    # 原始分布
    ax1.hist(df['礼物收入(美元)'], bins=50, color='#2E86AB', edgecolor='white')
    ax1.set_xlabel('礼物收入 ($)')
    ax1.set_ylabel('直播间数量')
    ax1.set_title('礼物收入分布')
    ax1.grid(True, alpha=0.3)

    # 对数分布
    log_income = np.log1p(df['礼物收入(美元)'])
    ax2.hist(log_income, bins=50, color='#A23B72', edgecolor='white')
    ax2.set_xlabel('log(1 + 礼物收入)')
    ax2.set_ylabel('直播间数量')
    ax2.set_title('礼物收入对数分布')
    ax2.grid(True, alpha=0.3)

    plt.tight_layout()
    plt.savefig(save_path, dpi=300, bbox_inches='tight')
    plt.close()


# ============================================================
# 七、生成Word报告
# ============================================================
def create_word_report(metrics_list, importance_df, df):
    """生成Word报告"""
    print("\n" + "=" * 80)
    print("正在生成Word分析报告...")
    print("=" * 80)

    doc = Document()
    title = doc.add_heading('问题3：直播间转化分析报告', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"报告生成时间: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph()

    # 一、分析目标
    doc.add_heading('一、分析目标', level=1)
    doc.add_paragraph(
        '本分析通过线性回归、Ridge回归和MLP神经网络回归模型，'
        '识别影响直播间礼物收入的关键特征，预测直播间收入，为流量分配和主播运营提供依据。'
    )

    # 二、数据概览
    doc.add_heading('二、数据概览', level=1)
    doc.add_paragraph(f"直播间样本数: {len(df):,}")
    doc.add_paragraph(f"平均礼物收入: ${df['礼物收入(美元)'].mean():.4f}")
    doc.add_paragraph(f"中位数礼物收入: ${df['礼物收入(美元)'].median():.4f}")
    doc.add_paragraph(f"高收入直播间比例: {df['是否高收入'].mean():.2%}")

    income_img = os.path.join(OUTPUT_DIR, '10_礼物收入分布.png')
    if os.path.exists(income_img):
        doc.add_picture(income_img, width=Inches(5.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 三、模型评估结果
    doc.add_heading('三、模型评估结果', level=1)
    table = doc.add_table(rows=1, cols=7)
    table.style = 'Light Grid Accent 1'
    hdr_cells = table.rows[0].cells
    headers = ['模型', '测试集RMSE', '测试集MAE', '测试集R²', '训练集RMSE', '训练集MAE', '训练集R²']
    for i, h in enumerate(headers):
        hdr_cells[i].text = h

    for metrics in metrics_list:
        row_cells = table.add_row().cells
        row_cells[0].text = metrics['model_name']
        row_cells[1].text = f"{metrics['test_rmse']:.4f}"
        row_cells[2].text = f"{metrics['test_mae']:.4f}"
        row_cells[3].text = f"{metrics['test_r2']:.4f}"
        row_cells[4].text = f"{metrics['train_rmse']:.4f}"
        row_cells[5].text = f"{metrics['train_mae']:.4f}"
        row_cells[6].text = f"{metrics['train_r2']:.4f}"

    doc.add_paragraph()

    # 四、真实值vs预测值
    doc.add_heading('四、真实值 vs 预测值', level=1)
    pred_img = os.path.join(OUTPUT_DIR, '10_真实值vs预测值.png')
    if os.path.exists(pred_img):
        doc.add_picture(pred_img, width=Inches(5.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 五、残差分析
    doc.add_heading('五、残差分析', level=1)
    for metrics in metrics_list:
        doc.add_paragraph(f"{metrics['model_name']} 残差图:")
        residual_img = os.path.join(OUTPUT_DIR, f'10_{metrics["model_name"]}_残差图.png')
        if os.path.exists(residual_img):
            doc.add_picture(residual_img, width=Inches(5.0))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 六、特征重要性
    doc.add_heading('六、线性回归特征重要性', level=1)
    doc.add_paragraph('基于线性回归系数绝对值排序的Top 15重要特征：')

    top15 = importance_df.head(15)
    table2 = doc.add_table(rows=1, cols=3)
    table2.style = 'Light Grid Accent 1'
    hdr_cells2 = table2.rows[0].cells
    for i, h in enumerate(['排名', '特征', '重要性']):
        hdr_cells2[i].text = h

    for idx, (_, row) in enumerate(top15.iterrows(), 1):
        row_cells = table2.add_row().cells
        row_cells[0].text = str(idx)
        row_cells[1].text = str(row['特征'])
        row_cells[2].text = f"{row['重要性']:.4f}"

    doc.add_paragraph()
    importance_img = os.path.join(OUTPUT_DIR, '10_线性回归_特征重要性.png')
    if os.path.exists(importance_img):
        doc.add_picture(importance_img, width=Inches(5.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 七、结论
    doc.add_heading('七、结论与建议', level=1)
    best_model = max(metrics_list, key=lambda x: x['test_r2'])
    doc.add_paragraph(f"1. 最优收入预测模型为 {best_model['model_name']}，测试集 R² 为 {best_model['test_r2']:.4f}。")

    if not importance_df.empty:
        top1 = importance_df.iloc[0]
        doc.add_paragraph(f"2. 最重要的收入驱动特征: {top1['特征']} (重要性: {top1['重要性']:.4f})。")

    doc.add_paragraph(
        '3. 建议平台优先扶持互动率高、观看人数多、主播等级高的直播间，'
        '并通过PK等玩法提升用户互动和付费转化。'
    )

    doc.save(DOC_FILE)
    print(f"Word报告已保存: {DOC_FILE}")


# ============================================================
# 八、主程序
# ============================================================
def main():
    sys.stdout = Tee(LOG_FILE)

    print("\n" + "=" * 80)
    print("问题3：直播间转化分析")
    print("=" * 80)
    print(f"开始时间: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    # 1. 加载与特征工程
    df = load_and_engineer_features()

    # 2. 准备建模数据
    X, y, numeric_features, categorical_features, df_model = prepare_modeling_data(df)
    X_train, X_test, y_train, y_test = train_test_split(
        X, y, test_size=0.2, random_state=42
    )
    print(f"\n训练集样本数: {len(X_train):,}")
    print(f"测试集样本数: {len(X_test):,}")

    preprocessor = build_preprocessor(numeric_features, categorical_features)

    # 3. 线性回归
    lr_model, lr_metrics = train_linear_regression(
        X_train, X_test, y_train, y_test, preprocessor
    )

    # 4. Ridge回归
    ridge_model, ridge_metrics = train_ridge(
        X_train, X_test, y_train, y_test, preprocessor
    )

    # 5. MLP回归
    mlp_model, mlp_metrics = train_mlp_regression(
        X_train, X_test, y_train, y_test, preprocessor
    )

    # 6. 特征重要性（基于线性回归）
    importance_df = get_feature_importance(lr_model, preprocessor,
                                           numeric_features, categorical_features)
    print("\n线性回归 Top 10 特征重要性:")
    print(importance_df.head(10).to_string(index=False))

    # 7. 可视化
    metrics_list = [lr_metrics, ridge_metrics, mlp_metrics]
    plot_feature_importance(importance_df, os.path.join(OUTPUT_DIR, '10_线性回归_特征重要性.png'))
    plot_actual_vs_predicted(metrics_list, os.path.join(OUTPUT_DIR, '10_真实值vs预测值.png'))
    plot_residuals(lr_metrics, os.path.join(OUTPUT_DIR, '10_线性回归_残差图.png'))
    plot_residuals(mlp_metrics, os.path.join(OUTPUT_DIR, '10_MLP神经网络回归_残差图.png'))
    plot_income_distribution(df, os.path.join(OUTPUT_DIR, '10_礼物收入分布.png'))

    # 8. Word报告
    create_word_report(metrics_list, importance_df, df)

    print("\n" + "=" * 80)
    print("问题3分析完成！")
    print(f"结束时间: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print("=" * 80)


if __name__ == '__main__':
    main()
